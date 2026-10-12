"""会议纪要整理:提示词 + 结构化解析 + **防虚构守卫** + docx 渲染。

【防虚构守卫为什么必须存在】
  大模型整理会议时最常见的失败不是漏信息,而是**把没说的补齐**:给待办编一个责任人、
  给会议填一个时间。周报里编一句摘要还能回看,导出的 .docx 是要提交归档的正式材料,
  错了就是台账造假。
  做法与本项目的「引用越界剥离」同构(答案里出现检索不到的 [n] 就删掉编号并记 badcase):
  **对可事实核查的字段**(会议时间、参会人员、待办的责任人与截止时间),
  逐项要求在**原始内容**里溯源;溯不到源的置为占位串,并把说明收进 `untraced`
  回给前端提示"请核对"。讨论要点/决议事项是改写而非引用,不做逐句溯源(会把正常的
  书面化改写全判成虚构),只做"空则占位"。

【占位串而不是留空】
  规格要求"原始信息缺失时保留占位提示"。留空在导出稿里看起来像"没写",
  占位串才是在提醒补录 —— 且占位串**不由模型书写**,避免它把占位写成各种花样。
"""
from __future__ import annotations

import io
import json
import re
from typing import Any

MINUTES_SECTIONS = ("会议主题", "会议时间", "参会人员", "讨论要点", "决议事项", "待办任务")

# 缺失占位串(与谈心谈话记录的两条占位分开设:这里语境是"纪要字段待补",不是"记录没填")
MISSING = "（未提供，请补充）"

MINUTES_PROMPT = """你是高校行政会议的纪要整理助手。把下面零散的会议原始内容整理成规范会议纪要。

输出**一个 JSON 对象**,键固定为:
{
  "title": "会议主题",
  "meeting_time": "会议时间",
  "attendees": "参会人员,顿号分隔",
  "discussion": ["讨论要点1", "讨论要点2"],
  "resolutions": ["决议事项1", "决议事项2"],
  "todos": [{"item": "待办事项", "owner": "责任人", "due": "截止时间"}]
}

硬性规则:
1. **只依据原始内容**。原始内容里没出现的姓名、时间、数字、结论,一律不要写;
2. 某个字段在原始内容里找不到依据 → 该值输出**空字符串**(数组输出空数组),
   不要自己编一个,也不要写"未提及""暂无"之类的文字(占位提示由系统统一加);
3. `meeting_time` 只写原文出现过的时间表述,原样保留(如"10月8日下午""2026-10-08 14:00");
4. `attendees` 只列原文提到的人或单位,不改名、不补职务;
5. 每条待办的 owner 与 due 必须是原文里明确说过的;原文只说了要做什么、没说谁做/何时做完,
   对应字段就留空字符串;
6. `discussion` 与 `resolutions` 用陈述句分条,保留原文中的数字、期限、条款号、金额;
7. 只输出 JSON,不要 markdown 代码围栏、不要解释文字。

会议原始内容:
{raw}
"""


# ---------------------------------------------------------------- 解析

def parse_minutes_output(raw: object) -> dict:
    """模型输出 → 纪要 dict。取不到 JSON 对象则抛 ValueError(调用方转 502)。"""
    if isinstance(raw, dict):
        return normalize_fields(raw)
    text = str(raw or "").strip()
    if not text:
        raise ValueError("整理结果为空")
    # 剥掉可能的 ```json 围栏,再取第一个平衡的 { ... }
    text = re.sub(r"^```[a-zA-Z]*\s*|^```\s*$", "", text).strip()
    start = text.find("{")
    if start < 0:
        raise ValueError("整理结果里没有找到 JSON")
    depth = 0
    for i in range(start, len(text)):
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
            if depth == 0:
                try:
                    # strict=False:容忍模型在字符串里吐**裸换行**(实测偶发),
                    # 否则一个控制字符就让整份纪要解析失败。
                    return normalize_fields(json.loads(text[start:i + 1], strict=False))
                except json.JSONDecodeError as e:
                    raise ValueError(f"整理结果 JSON 解析失败:{e}") from e
    raise ValueError("整理结果的 JSON 括号不平衡")


def _clean_item(s: object) -> str:
    """去掉条目自带的编号/项目符号(模型给「1.了解…」或「- 了解…」时不再叠加一层符号)。"""
    t = re.sub(r"^\s*(?:[-·•*]\s*)?(?:\d{1,2}\s*[.、)）]\s*)?", "", str(s or ""))
    return t.strip(" \t-·•*")


def _as_lines(v: Any) -> list[str]:
    """数组 / 换行文本 → 条目列表(模型两种形态都可能给)。"""
    if v is None:
        return []
    if isinstance(v, str):
        return [c for c in (_clean_item(p) for p in re.split(r"[\n;；]+", v)) if c]
    if isinstance(v, list):
        out: list[str] = []
        for item in v:
            if isinstance(item, dict):
                item = item.get("item") or item.get("text") or item.get("content") or ""
            c = _clean_item(item)
            if c:
                out.append(c)
        return out
    c = _clean_item(v)
    return [c] if c else []


def _as_todos(v: Any) -> list[dict]:
    out: list[dict] = []
    if isinstance(v, str):
        v = _as_lines(v)
    if not isinstance(v, list):
        return out
    for item in v:
        if isinstance(item, dict):
            todo = {
                "item": _clean_item(item.get("item") or item.get("task")
                                    or item.get("content") or ""),
                "owner": str(item.get("owner") or item.get("person") or "").strip(),
                "due": str(item.get("due") or item.get("deadline") or "").strip(),
            }
        else:
            todo = {"item": _clean_item(item), "owner": "", "due": ""}
        if todo["item"]:
            out.append(todo)
    return out


def _first_text(obj: dict, *names: str) -> str:
    """取第一个**非空字符串**字段(不把 list/dict `str()` 成 "['a']" 塞进纪要)。"""
    for n in names:
        v = obj.get(n)
        if isinstance(v, str) and v.strip():
            return v.strip()
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return str(v)
    return ""


def normalize_fields(obj: dict) -> dict:
    """补齐六个字段的标准形态(键恒在、类型恒对),供守卫与渲染直接消费。"""
    if not isinstance(obj, dict):
        obj = {}
    return {
        "title": _first_text(obj, "title", "会议主题"),
        "meeting_time": _first_text(obj, "meeting_time", "time", "会议时间"),
        "attendees": _first_text(obj, "attendees", "attendee", "参会人员", "参会人"),
        "discussion": _as_lines(obj.get("discussion") or obj.get("讨论要点")),
        "resolutions": _as_lines(obj.get("resolutions") or obj.get("决议事项")),
        "todos": _as_todos(obj.get("todos") or obj.get("待办任务")),
    }


# ---------------------------------------------------------------- 防虚构守卫

_NORM_RE = re.compile(r"[\s\u3000]+")
_DIGITS_RE = re.compile(r"\d+")
_SPLIT_RE = re.compile(r"[、，,;；/]| 和 | 与 ")


def _norm(s: object) -> str:
    return _NORM_RE.sub("", str(s or ""))


def _digit_runs(s: object) -> list[str]:
    """数字碎片归一(去前导零):原文写「8日」、模型写「08」不该因此判成虚构。"""
    return [str(int(d)) for d in _DIGITS_RE.findall(str(s or ""))]


def traceable(value: str, src_norm: str) -> bool:
    """时间类字段能否溯源:整串命中原文,或**其中每个数字都出现在原文里**。

    判得严是有意的:纪要里的时间要进台账,模型补出一个原文没有的年份,与漏写一条要点
    不是一个量级的问题。代价是模型把「10月8日下午」重排成「2026-10-08 14:00」时会判不过、
    回落占位 —— 界面提示核对,辅导员补一下即可。
    """
    v = _norm(value)
    if not v:
        return False
    if v in src_norm:
        return True
    runs = _digit_runs(v)
    if not runs:
        return False
    return all(r in _digit_runs(src_norm) for r in runs)


def split_attendees(s: str) -> list[str]:
    return [p.strip() for p in _SPLIT_RE.split(_norm(s)) if p.strip()]


def guard(minutes: dict, source_text: str) -> tuple[dict, list[str]]:
    """对可核查字段逐项溯源;溯不到 → 置占位串并记入 untraced 说明。

    先跑一遍 `normalize_fields`,所以调用方可以直接把模型输出的原始 dict 丢进来。
    讨论要点/决议事项为空、待办为空都**不报错**,按占位处理 ——
    "这次会没形成决议"和"模型没读出来"在原文里无法区分,交给辅导员核对。
    """
    out = normalize_fields(minutes)
    untraced: list[str] = []
    src_norm = _norm(source_text)

    if not out["title"]:
        out["title"] = MISSING

    if not out["meeting_time"]:
        out["meeting_time"] = MISSING
        untraced.append("会议时间")
    elif not traceable(out["meeting_time"], src_norm):
        out["meeting_time"] = MISSING
        untraced.append("会议时间（整理结果未能在原文中溯源）")

    names = split_attendees(out["attendees"])
    if not names:
        out["attendees"] = MISSING
        untraced.append("参会人员")
    else:
        kept = [n for n in names if n in src_norm]
        dropped = [n for n in names if n not in kept]
        if not kept:
            out["attendees"] = MISSING
            untraced.append("参会人员（整理结果未能在原文中溯源）")
        else:
            out["attendees"] = "、".join(kept)
            if dropped:
                untraced.append(f"参会人员（{'、'.join(dropped)} 未能在原文中溯源，已剔除）")

    for key, label in (("discussion", "讨论要点"), ("resolutions", "决议事项")):
        if not out[key]:
            out[key] = []
            untraced.append(label)

    for i, t in enumerate(out["todos"], 1):
        owner, due = t["owner"], t["due"]
        owner_ok = bool(owner) and _norm(owner) in src_norm
        due_ok = bool(due) and _norm(due) in src_norm
        t["owner"] = owner if owner_ok else MISSING
        t["due"] = due if due_ok else MISSING
        if owner and not owner_ok:
            untraced.append(f"待办{i}责任人（未能在原文中溯源）")
        if due and not due_ok:
            untraced.append(f"待办{i}截止时间（未能在原文中溯源）")
    if not out["todos"]:
        untraced.append("待办任务")
    return out, untraced


# ---------------------------------------------------------------- 渲染

def render_markdown(m: dict) -> str:
    """纪要 → Markdown(与既有周报 content_md 同构)。"""
    lines = [f"# {m.get('title') or MISSING}", ""]
    lines += [f"- **会议时间**：{m.get('meeting_time') or MISSING}",
              f"- **参会人员**：{m.get('attendees') or MISSING}", ""]
    for label, key in (("讨论要点", "discussion"), ("决议事项", "resolutions")):
        lines.append(f"## {label}")
        items = m.get(key) or []
        lines += [f"- {x}" for x in items] or [f"- {MISSING}"]
        lines.append("")
    lines.append("## 待办任务")
    todos = m.get("todos") or []
    if not todos:
        lines.append(f"- {MISSING}")
    else:
        lines += ["| 事项 | 责任人 | 截止时间 |", "| --- | --- | --- |"]
        lines += [f"| {t.get('item', '')} | {t.get('owner') or MISSING} "
                  f"| {t.get('due') or MISSING} |" for t in todos]
    lines += ["", "> 本草稿由智枢 Agent 基于会议原始内容整理,未能溯源的字段已标为占位,"
              "请在归档前核对补全。"]
    return "\n".join(lines)


def render_docx(m: dict) -> bytes:
    """纪要 → .docx 字节。

    中文默认字体显式设 `w:eastAsia`:只设 `font.name` 在部分 Word 版本下中文会回退,
    轻则字形不统一,重则显示成方框 —— 而这份文件是要交上去的。
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    rfonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "宋体")

    doc.add_heading(m.get("title") or MISSING, level=1)

    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for label, value in (("会议时间", m.get("meeting_time") or MISSING),
                         ("参会人员", m.get("attendees") or MISSING)):
        cells = info.add_row().cells
        cells[0].text, cells[1].text = label, value

    for label, key in (("讨论要点", "discussion"), ("决议事项", "resolutions")):
        doc.add_heading(label, level=2)
        items = m.get(key) or []
        if items:
            for x in items:
                doc.add_paragraph(x, style="List Bullet")
        else:
            doc.add_paragraph(MISSING)

    doc.add_heading("待办任务", level=2)
    todos = m.get("todos") or []
    if todos:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ("事项", "责任人", "截止时间")):
            cell.text = text
        for t in todos:
            cells = table.add_row().cells
            cells[0].text = t.get("item", "")
            cells[1].text = t.get("owner") or MISSING
            cells[2].text = t.get("due") or MISSING
    else:
        doc.add_paragraph(MISSING)

    doc.add_paragraph("")
    doc.add_paragraph("本纪要由智枢 Agent 基于会议原始内容整理；标注占位的字段未能在原文中溯源，"
                      "请在归档前核对补全。")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
